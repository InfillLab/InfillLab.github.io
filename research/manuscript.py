from pathlib import Path
from copy import deepcopy
import json,re,sys
from docx import Document
from docx.shared import Pt,Cm,RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT,WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

src=Path(sys.argv[1] if len(sys.argv)>1 else 'C:/Users/mihae/Downloads/Engineering_Review_Transition_Material_Interlocking_Authors_References.docx')
out=Path(sys.argv[2] if len(sys.argv)>2 else 'C:/Users/mihae/Downloads/Engineering_Review_Transition_Material_Interlocking_Results_Revised.docx')
results=Path(sys.argv[3] if len(sys.argv)>3 else 'C:/Users/mihae/Downloads/InfillLab_Research_Results')
stats=json.loads((results/'statistics.json').read_text());combined=json.loads((results/'combined_statistics.json').read_text())
d=Document(src);body=d.element.body
refs=[deepcopy(p._p) for p in d.paragraphs if re.match(r'^\[\d+\]',p.text)]
old_discussion_table=deepcopy(d.tables[-1]._tbl)

def replace(prefix,text):
 p=next(p for p in d.paragraphs if p.text.startswith(prefix));p.text=text
 for r in p.runs:r.font.name='Times New Roman';r.font.size=Pt(11)

abstract=('Material switching through a shared melt path generates a transition discharge whose reuse depends on composition, receiving capacity and deposition order. This study formulates a constrained allocation model and evaluates a seeded adaptive scheduler against reference-order and greedy schedules. Ten simplified receiving patterns are compared in 360 main configurations, followed by composition and segmentation sensitivity analyses. In total, 540 distinct configurations are evaluated using 30 seeds and 200 search iterations, giving 16,200 adaptive runs. Conservation provides an upper utilisation bound, while an exact two-segment example shows that the shortest geometric route can violate composition constraints. In the main comparison, the median paired utilisation gain over the better baseline is zero; positive median gains occur in 34 of 360 configurations. The largest gain is 47.73 percentage points for a concentric pattern. Six pattern families have fixed execution order and provide no search benefit. Refining segment length from 1 to 0.5 mm changes individual utilisation values by -0.77 to 0.26 percentage points and substantially increases computation time. Independent schedule checks identify no volume-accounting or composition-window violations. The results support composition-aware scheduling for selected receiving geometries, but do not establish a universally superior infill or increased joint strength. Physical transition calibration and matched mechanical tests remain necessary before structural use.')
assert len(abstract.split())<=250
cell=d.tables[0].cell(1,2);cell.text=abstract
for p in cell.paragraphs:
 p.paragraph_format.first_line_indent=Cm(0);p.paragraph_format.space_after=Pt(0)
 for r in p.runs:r.font.name='Times New Roman';r.font.size=Pt(10)
replace('The contribution is a theoretical formulation','The contribution is a formulation and computational evaluation of that coupling. An analytical utilisation bound and an exact routing counterexample are followed by controlled comparisons of ten receiving patterns. A seeded adaptive scheduler is compared with two deterministic baselines under the same geometry and acceptance constraints. The study asks where additional search improves transition reuse, how this improvement depends on path order, and how composition and segmentation assumptions affect the result. Reinforcement is not inferred from reduced disposal.')
replace('For segment ','For segment i, the required nominal volume Vi is calculated from its centreline length li and constant bead area Ai. The nominal receiving capacity Vcap is the sum of these path volumes. In the computational study, Ai is bead width multiplied by layer height. Equation (1) is a path-volume budget; it is not a Boolean union of deposited solids, and crossings or adjacent bead overlaps are not subtracted.')
replace('A rectangular bead approximation','This rectangular bead approximation permits reproducible allocation comparisons, but does not resolve the true bead profile or local overfilling. Sharp turns, intersections and occupied material can reduce physical receiving capacity. Unassigned segments remain unfilled in the simulation. Completing them with clean material would require an additional manufacturing schedule, which is outside the reported transition event.')
replace('The interval is partitioned','The interval can be partitioned into volume increments ΔVj for accounting. An eligibility indicator ej equals one when an increment satisfies the assumed global composition window and zero otherwise. An allocation indicator yj equals one when that increment is deposited into an admissible receiving segment. Eligibility and allocation are distinct: globally eligible material can be incompatible with the next accessible local segment. The numerical scheduler accepts whole segments by checking both ends of their discharge intervals under a monotonic composition history.')
replace('In Eq. (2),','In Eq. (2), Velig is the globally eligible transition volume and Vuse is the volume assigned to receiving segments. Increments crossing a boundary must be split for exact accounting. The implementation obtains these boundaries by analytical inversion of the composition curve; it does not approximate eligibility by a sampled mean. Segment length remains a separate allocation discretisation, examined in Section 6.4. All unassigned transition volume is counted as external discard within this model.')
replace('Each segment has two possible','A schedule specifies an ordered subset of accepted segments and the permitted traversal direction of each segment. Reversible segments have two choices; connected or monotonic chains retain their direction and predecessor constraints. Let si and ti be the start and end points under the selected direction, and let s0 be the initial nozzle position. Equation (5) defines non-deposition travel for the n accepted segments. Distances are Euclidean in the simulated layer, with no final return move. Nozzle collisions and obstacle-avoiding travel are not modelled.')
replace('Every required segment must','An accepted segment is deposited once. A complete schedule assigns every required segment; a partial schedule reports its remaining unassigned segments explicitly. Cumulative extrusion includes any purge needed to reach an admissible composition. The discharge cursor cannot move backwards. The implemented constraints preserve chain order and traversal direction where specified; they do not establish physical support, collision clearance or printability.')
replace('For an initial material-efficiency','Equation (6) defines a dimensionless score from discarded fraction and normalised travel. In the computational study, Lref is fixed at 100 mm for every geometry and segmentation, and α = 0.65. Ranking is lexicographic: complete schedules precede partial schedules, greater receiver coverage precedes lower coverage, and J breaks remaining ties. Coverage is Vuse/Vcap. Consequently, the selected weight does not trade away coverage in favour of shorter travel; the score is a secondary routing criterion.')
replace('Small instances admit','Small instances admit exhaustive enumeration and provide reference solutions. For larger instances, adaptive search reinforces segment-to-segment connections in its best schedule while sampling alternative feasible connections. This supplies a conceptual link to slime-mould-inspired transport adaptation. The implemented method is an experimental reinforcement heuristic, not the canonical Slime Mould Algorithm, and no biological material is used.')
replace('A candidate bioinspired solver','Candidate construction selects an admissible next segment and direction, advances the discharge cursor and accounts for preceding purge. It stops when no admissible candidate remains. There is no repair stage or look-ahead guarantee that the remaining receiver can be completed. The precise baseline, sampling, reinforcement and termination rules used in the experiments are specified in Section 5.3.')
cut=next(p._p for p in d.paragraphs if p.text.startswith('5 Discussion'));active=False
for el in list(body):
 if el is cut:active=True
 if active and el.tag!=qn('w:sectPr'):body.remove(el)
normal=d.styles['Normal'];normal.font.name='Times New Roman';normal.font.size=Pt(11);normal.font.color.rgb=RGBColor(0,0,0)
normal.paragraph_format.line_spacing=1;normal.paragraph_format.first_line_indent=Cm(.6);normal.paragraph_format.space_after=Pt(0)

def para(text,kind='body'):
 p=d.add_paragraph(text);f=p.paragraph_format;f.line_spacing=1;f.space_after=Pt(0);f.widow_control=True;f.first_line_indent=Cm(.6 if kind=='body' else 0)
 for r in p.runs:r.font.name='Times New Roman';r.font.size=Pt(11);r.font.color.rgb=RGBColor(0,0,0)
 if kind in ('h1','h2'):
  f.keep_with_next=True;f.space_before=Pt(11 if kind=='h1' else 8);f.space_after=Pt(5)
  for r in p.runs:r.bold=kind=='h1';r.italic=kind=='h2';r.font.size=Pt(12 if kind=='h1' else 11)
 if kind=='caption':
  f.space_before=Pt(5);f.space_after=Pt(5)
  for r in p.runs:r.italic=True
 return p

def table(caption,headers,rows,widths):
 para(caption,'caption').paragraph_format.keep_with_next=True
 t=d.add_table(rows=1,cols=len(headers));t.autofit=False;t.alignment=WD_TABLE_ALIGNMENT.CENTER
 for col,width in zip(t.columns,widths):col.width=Cm(width)
 for c,text,width in zip(t.rows[0].cells,headers,widths):c.text=text;c.width=Cm(width)
 for row in rows:
  for c,text,width in zip(t.add_row().cells,row,widths):c.text=str(text);c.width=Cm(width)
 pr=t._tbl.tblPr;borders=OxmlElement('w:tblBorders')
 for edge in ['top','left','bottom','right','insideH','insideV']:
  e=OxmlElement('w:'+edge);e.set(qn('w:val'),'single' if edge in ('top','bottom') else 'nil');e.set(qn('w:sz'),'6');e.set(qn('w:color'),'000000');borders.append(e)
 pr.append(borders);margins=OxmlElement('w:tblCellMar')
 for edge in ['top','bottom','left','right']:
  e=OxmlElement('w:'+edge);e.set(qn('w:w'),'55' if edge in ['top','bottom'] else '75');e.set(qn('w:type'),'dxa');margins.append(e)
 pr.append(margins)
 for ir,row in enumerate(t.rows):
  trpr=row._tr.get_or_add_trPr();trpr.append(OxmlElement('w:cantSplit'))
  if ir==0:trpr.append(OxmlElement('w:tblHeader'))
  for ic,c in enumerate(row.cells):
   c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
   if ir==0:
    cb=OxmlElement('w:tcBorders');e=OxmlElement('w:bottom');e.set(qn('w:val'),'single');e.set(qn('w:sz'),'4');cb.append(e);c._tc.get_or_add_tcPr().append(cb)
   for p in c.paragraphs:
    p.paragraph_format.first_line_indent=Cm(0);p.paragraph_format.space_after=Pt(0);p.paragraph_format.line_spacing=1;p.paragraph_format.keep_with_next=ir==0;p.alignment=WD_ALIGN_PARAGRAPH.LEFT if ic==0 else WD_ALIGN_PARAGRAPH.CENTER
    for r in p.runs:r.font.name='Times New Roman';r.font.size=Pt(11);r.bold=ir==0
 return t

def figure(name,caption):
 p=d.add_paragraph();p.paragraph_format.first_line_indent=Cm(0);p.paragraph_format.space_before=Pt(6);p.paragraph_format.space_after=Pt(0);p.paragraph_format.keep_with_next=True;p.alignment=WD_ALIGN_PARAGRAPH.CENTER
 p.add_run().add_picture(str(results/'figures'/name),width=Cm(16.891));para(caption,'caption')

para('5 Computational investigation','h1')
para('5.1 Geometry and composition assumptions','h2')
para('All receiving geometries occupy one nominal 40 mm × 18 mm layer. Bead width is 0.5 mm, layer height is 0.2 mm and path spacing is 2 mm. The initial nozzle position is the lower-left origin of the layer. The nine reference names are Concentric, Rectilinear, Monotonic, Monotonic line, Global monotonic line, Aligned rectilinear, Hilbert curve, Archimedean chords and Octagram spiral. These are simplified centreline constructions, not exact reproductions of a commercial slicer. The tenth pattern, termed transition-aware interlaced infill (TAII), consists of staggered finger paths. Its name denotes the proposed receiving construction and allocation approach, not verified three-dimensional locking or mechanical superiority.')
para('Concentric loops retain their internal order, but different loops can be selected at different stages. Rectilinear, Aligned rectilinear and TAII segments can be reordered and reversed. The three monotonic variants retain global sequence and direction. The Hilbert and spiral curves retain their connected sequence and direction. Accordingly, six families provide no alternative deposition order to the search. Monotonic line and Global monotonic line coincide on this single rectangular region; Rectilinear and Aligned rectilinear share geometry but differ in reference traversal directions. The ten labels therefore do not represent ten independent topologies.')
para('Native mode retains each generated path set. Nominal capacities range from 19.526 mm³ for Archimedean chords to 35.200 mm³ for Concentric. Equal-capacity mode scales the retained length of every native path by the same pattern-specific factor, trimming its end so that each pattern has 19.526 mm³ total capacity. This controls the nominal material budget but changes some path shapes and spatial coverage. It does not impose equal stiffness, connectivity or occupied area. Nominal allocation segments are 2 mm long in the main experiment, with shorter residual segments at path ends.')
para('The assumed fraction of material B is c(v) = (v/Vtr)^γ for switching from A to B and c(v) = 1 - (v/Vtr)^γ for the reverse direction. The global admissible interval is [0.1, 0.9]. Its eligible fraction is 0.9^(1/γ) - 0.1^(1/γ) for either direction under these symmetric thresholds. Thus the eligible fraction is 0.80 for γ = 0.5 or 1, and 0.63246 for γ = 2. These power-law curves and thresholds are prescribed sensitivity inputs, not fits to PLA-TPU discharge measurements.')
para('For graded acceptance, the target composition of segment i is 0.15 + 0.70 times its mean vertex y-coordinate divided by the region height. Its local window is this target plus or minus δ, clipped to [0, 1]. Uniform acceptance uses [0, 1] locally, while the global [0.1, 0.9] restriction remains active. The intersection of global and local windows is inverted to obtain an admissible discharge interval. A segment is placed at the earliest admissible volume after the current cursor only if its entire volume fits. Any intervening discharge and the unused tail are counted as external discard.')
para('5.2 Experimental design','h2')
para('Table 3 defines the configuration matrix. The main experiment crosses ten patterns, two capacity modes, three ratios r = Vtr/Vcap, three local-window settings and two switching directions. The nonlinear and segmentation studies use equal capacity and r = 1 to avoid introducing a simultaneous capacity change. Twenty 2 mm configurations are shared with the main experiment and are counted only once. A further 0.5 mm segmentation check was added after inspection of the 1 mm and 2 mm results; it is reported as a follow-up sensitivity check.')
table('Table 3. Configuration matrix and computation budget',['Study','Factor settings','Configurations'],[
 ['Main comparison','10 patterns × 2 capacities × 3 ratios × 3 windows × 2 directions','360'],
 ['Composition sensitivity','10 patterns × 2 graded windows × 2 directions × 2 additional exponents','80 new'],
 ['Segmentation sensitivity','10 patterns × 2 directions × 5 segment lengths','100, including 20 reused'],
 ['Follow-up segmentation','10 patterns × 2 directions × 0.5 mm','20 new'],
 ['Total unique','30 seeds and 200 search iterations per configuration','540']],[3.9,9.7,3.3])
para('The main ratios are 0.5, 1 and 2. The three local-window settings are graded δ = 0.10, graded δ = 0.25 and uniform acceptance. The main composition exponent is γ = 1; additional values are 0.5 and 2. Segmentation lengths are 1, 2, 4, 6 and 8 mm, followed by 0.5 mm, with δ = 0.10. Every unique configuration uses seeds 1 to 30 and 200 candidate-construction iterations per seed. This gives 16,200 adaptive runs. Two deterministic schedules are also saved once per configuration, giving 1,080 baseline results and 17,280 stored schedule results in total.')
para('5.3 Scheduling methods','h2')
para('The reference-order baseline attempts segments in generated order and forward direction, skipping a segment if it cannot be placed. The greedy baseline, labelled nearest feasible in the software, repeatedly selects an available segment and direction with minimum travel distance plus 3 mm/mm³ multiplied by preceding purge volume. It therefore penalises both travel and purge rather than minimising distance alone. Both baselines obey the same acceptance and predecessor constraints as adaptive search. Each configuration is compared with the better of its two baselines according to the ranking in Section 3.4.')
para('Adaptive search starts from that better baseline. Each candidate schedule is constructed from an initially empty receiver. Available placements are sorted by the same travel-plus-purge cost q. With probability 0.25, the cheapest placement is selected. Otherwise, selection is weighted over the cheapest 65% of placements, rounded upwards, with at least four placements included when available. A placement weight equals its current directed-edge preference divided by 1 + q, where q is the numerical cost expressed in millimetres and the added unit corresponds to 1 mm. The pseudo-random generator is Mulberry32 with the declared integer seed.')
para('All edge preferences initially equal one. After each candidate schedule, the incumbent is replaced only if the candidate ranks better. Every preference p is relaxed to 0.98p + 0.02, and each edge in the retained incumbent receives an increment of 0.12, capped at 12. The edge index records the previous segment and next segment, including a start node; it does not distinguish traversal direction. Search stops after 200 iterations. A run therefore evaluates 200 sampled schedules and two initial baselines, or 202 schedules. There is no matching 202-evaluation budget for the deterministic baselines, whose role is to establish whether extra search provides a benefit.')
para('Since the better baseline is retained, a non-worse final rank is guaranteed by construction. It is not independent evidence of optimiser quality. A gain must exceed that baseline. The study does not include an ablation without reinforcement, a conventional local-search comparator or an exact optimum for the larger geometries, so it cannot isolate the contribution of reinforcement or claim global optimality.')
para('5.4 Verification and data reduction','h2')
para('Before aggregation, saved schedules were checked for duplicate placement, predecessor order, fixed-direction constraints, chronological discharge, segment volume, local and global composition windows, travel accounting, total volume balance, coverage and the analytical upper bound. An independent Python pass recomputed volume sums, composition endpoints, precedence, coverage and completion from the saved geometry and steps. All 17,280 results passed. The largest absolute volume-balance residual was 7.11 × 10^-15 mm³, below the checking tolerance of 10^-7 mm³. The exact two-segment result reproduced 28.284271 mm feasible travel and full allocation of 12 mm³.')
para('The implementation also passed 21 automated checks, including the analytical bound cases, reversed composition histories, seed repeatability, non-overlapping local and global windows, an empty receiver, insufficient transition volume and fixed travel normalisation across segment resolutions. Zero transition volume is excluded by input validation because utilisation would be undefined. These checks establish internal consistency under the implemented assumptions, not physical validation.')
para('For each configuration, the adaptive result is summarised by the median over 30 seeds. Comparisons are paired with the better deterministic baseline for that configuration. Group summaries then give equal weight to each configuration median and report the median and interquartile range. A positive utilisation improvement is counted above 10^-6 percentage points to suppress numerical round-off. These are descriptive summaries of the chosen design matrix; the seeds are algorithmic repeats, not independent specimens, and the interquartile ranges are not confidence intervals.')
para('Calculations used model version 2.1.0 under Node.js 24.19.0 on Windows with an Intel Core i7-9750H processor and three concurrent worker threads. Run times include adaptive initialisation, the two internal baselines and 200 search iterations. They are wall-clock computational times under concurrent execution, not isolated processor benchmarks or printing-time predictions. The accompanying research archive contains configuration manifests, model source hashes, complete compressed schedules, seed traces, tabulated outputs and the analysis and plotting scripts.')
para('6 Computational results','h1').paragraph_format.page_break_before=True
para('6.1 Pattern-dependent benefit of adaptive scheduling','h2')
para('Across the 360 main configurations, the median paired utilisation gain was 0.00 percentage points, with an interquartile range of 0.00 to 0.00. Positive median utilisation gains occurred in 34 configurations (9.44%); 326 were tied within the numerical threshold. Median solution rank improved in 43 configurations (11.94%), including nine cases with unchanged utilisation and lower travel. Table 4 and Figure 1 show that the benefit was concentrated in geometries with alternative segment choices.')
names=[('concentric','Concentric'),('rectilinear','Rectilinear'),('monotonic','Monotonic'),('monotonic-line','Monotonic line'),('global-monotonic','Global monotonic line'),('aligned','Aligned rectilinear'),('hilbert','Hilbert curve'),('archimedean','Archimedean chords'),('octagram','Octagram spiral'),('taii','TAII')]
rows=[]
for pid,name in names:
 g=stats['pattern'][pid];rows.append([name,str(g['improved_use']),str(g['improved_rank']),f"{g['gain_pp']['median']:.2f}",f"{g['gain_pp']['max']:.2f}"])
table('Table 4. Main comparisons by pattern, with 36 configurations per pattern',['Pattern','Higher utilisation','Better rank','Median gain, percentage points','Largest gain, percentage points'],rows,[4.8,2.2,2,3.9,4])
para('Concentric paths accounted for 19 utilisation improvements and 27 rank improvements. Their largest median gain was 47.73 percentage points for native capacity, r = 1, δ = 0.25 and A-to-B switching. Utilisation increased from 29.55% for the better baseline to 77.27% for the adaptive median, while receiving-region travel increased from 0.71 to 175.23 mm. This illustrates the coverage-first ranking: more material was assigned at a substantial travel cost. Rectilinear and Aligned rectilinear each improved utilisation in seven configurations, with maximum median gains of 1.64 and 1.74 percentage points. TAII improved utilisation in one configuration, by 0.51 percentage points, and rank in two. It was not a generally superior receiving pattern in this study.')
figure('Figure_1_pattern_gains.png','Figure 1. Utilisation gain over the better deterministic baseline in the 360 main configurations: (a) full range; (b) enlarged view from 0 to 2 percentage points. Open circles show configuration-level differences using the adaptive median over 30 seeds; filled diamonds show pattern medians. Vertical offsets separate overlapping points and carry no numerical meaning. Values above the range of panel (b) remain visible in panel (a).')
para('The three monotonic variants and the three fixed-order curves showed no gain. Their predecessor and direction constraints leave no scheduling alternatives, so extra randomised construction cannot change the executed order. The observed differences therefore combine geometric effects with the available scheduling freedom. The four families that allow branching account for 144 configurations; 34 of these (23.61%) improved median utilisation. This conditional proportion is more informative about the search opportunity than treating all ten names as equally reorderable.')
para('6.2 Transition volume and receiving coverage','h2')
para('Figure 2(a) separates utilisation of the transition stream from filling of the receiver. At r = 0.5, median adaptive utilisation was 77.69% (interquartile range 31.46% to 79.60%), while median receiver coverage was only 38.84%. At r = 1, these quantities coincided at 78.29% (32.46% to 79.83%). At r = 2, median utilisation decreased to 50.00% (30.12% to 50.00%), while median coverage reached 100.00%. Each ratio summarises 120 configuration medians. Lower utilisation at r = 2 is therefore partly a capacity consequence: even a filled receiver uses only half the available transition volume.')
para('With γ = 1 and the global [0.1, 0.9] restriction, the utilisation bounds are 80%, 80% and 50% at r = 0.5, 1 and 2, respectively. No adaptive run completed the receiver at r = 0.5 or r = 1, because the eligible stream is insufficient. At r = 2, complete schedules occurred in 59.06% of the 3,600 adaptive runs. The remaining schedules were valid partial allocations. Their unfilled volume cannot be interpreted as a completed printed part.')
para('6.3 Sensitivity to composition history','h2')
para('The matched composition comparison contains 40 configurations per exponent, using equal capacity, r = 1, two graded windows and both switching directions. Median utilisation was 37.15% for γ = 0.5, 43.63% for γ = 1 and 38.45% for γ = 2. Corresponding interquartile ranges were 19.31% to 65.88%, 31.01% to 78.02%, and 18.28% to 58.36%. Figure 2(b) shows that changing the composition curve changes attainable allocation as well as, for γ = 2, the aggregate eligibility bound. Equal aggregate eligibility at γ = 0.5 and γ = 1 does not imply equal compatibility with local windows.')
figure('Figure_2_capacity_and_composition.png','Figure 2. Utilisation under changes in (a) transition-volume ratio and (b) composition exponent. Points are medians of configuration medians; bars span the interquartile range, not a confidence interval. Panel (a) uses 120 configurations per ratio; panel (b) uses 40 matched configurations per exponent. Dashed lines show the analytical bound.')
para('6.4 Sensitivity to segmentation and computation time','h2')
para('The segmentation comparison contains 20 matched pattern-direction cases at each length. Table 5 reports utilisation and run time after taking the median over 30 seeds for each case. Although the pooled utilisation median is 32.46% at 0.5, 1, 2 and 4 mm, individual configurations change. Refinement from 2 to 1 mm changes utilisation by -2.05 to 1.02 percentage points. The follow-up refinement from 1 to 0.5 mm changes it by -0.77 to 0.26 percentage points, with zero median paired change.')
rows=[]
for length in [.5,1,2,4,6,8]:
 g=combined['segmentation'][str(float(length))];rows.append([f'{length:g}',f"{g['eta_median_pct']:.2f}",f"{g['eta_q25_pct']:.2f} - {g['eta_q75_pct']:.2f}",f"{g['runtime_median_ms']:.2f}"])
table('Table 5. Segmentation sensitivity across 20 matched configurations per length',['Segment length, mm','Median utilisation, %','Interquartile range, %','Median run time, ms'],rows,[4,4.3,4.3,4.3])
para('The median of configuration-level run times increased from 23.29 ms at 2 mm to 90.67 ms at 1 mm and 344.51 ms at 0.5 mm. These increases are consistent with the larger candidate set and preference matrix, although concurrent execution affects timing. Figure 3 shows paired utilisation changes relative to 2 mm alongside computational cost. The 0.5 mm runs were performed as a separate follow-up batch under the same worker-count setting.')
figure('Figure_3_segmentation.png','Figure 3. Segmentation sensitivity: (a) paired utilisation change relative to the 2 mm case; (b) run time on a logarithmic scale. Panel (a) boxes span the interquartile range, with median lines, whiskers to 1.5 interquartile ranges and outlying points. Panel (b) shows medians and interquartile ranges across 20 configuration medians.')
para('These data do not demonstrate strict discretisation convergence. Changing segment length also changes the segment vertices used to define local acceptance windows, while the search budget remains fixed. The comparison therefore changes both allocation granularity and aspects of the constraint field. A convergence study would need a continuous spatial acceptance field held fixed during refinement, together with evidence that search error is controlled. The present results quantify sensitivity of this particular implementation.')
para('7 Discussion','h1')
para('The calculations support a limited conclusion: composition-aware ordering can recover more admissible transition material in some geometries, particularly when several independently accessible paths compete for the same discharge interval. The large concentric improvement and the many zero-gain cases are both consequences of the selected geometry, composition field and scheduling constraints. They do not establish that Concentric is mechanically preferable, nor that adaptive scheduling is necessary for every pattern.')
para('The proposed TAII construction provides a testable receiver geometry but shows little allocation advantage in the evaluated matrix. Its mechanical value remains unresolved. A useful next comparison would vary finger dimensions and the continuous acceptance field while keeping nominal capacity and scheduling freedom comparable. The present data do not justify selecting it solely because it is the tenth pattern or because it resembles an interlock.')
para('Several limitations restrict transfer to manufacturing. Nominal volume ignores bead overlap and occupied solids; the single-layer model omits support, shell clearance, thermal history, pressure dynamics and travel to a purge location. Reported travel covers moves between accepted receiving segments only. It excludes motion to and from external discard locations and cannot quantify total printer travel or cycle time. Unfilled segments would require a feasible completion strategy. The composition curve and all acceptance windows require calibration for a specific shared-melt-path system. Separate-nozzle and resin processes are outside this model.')
para('Mechanical suitability cannot be obtained from the utilisation score. The geometry studies [5] - [7] justify testing interfacial design, while the transition measurements in [13] justify treating composition history as process-dependent. Neither establishes the properties of the assumed transition stream. Table 6 identifies the physical comparisons needed to distinguish geometry from allocation. These are proposed validation groups, not completed experiments.')
para('Table 6. Physical comparisons required before a reinforcement claim','caption').paragraph_format.keep_with_next=True
body.insert(len(body)-1,old_discussion_table)
para('The clean-material and transition-allocated interlocks should have matched geometry, actual deposited mass, build orientation and processing history. Tensile or peel loading should activate the intended separation mode, with force-displacement response, work to failure and fracture location recorded. Compression and absorbed energy require a common terminal strain and normalisation by measured mass. Replication should be chosen from pilot variability and a predeclared effect or non-inferiority margin. None of these physical responses is predicted by the present scheduler.')
para('Algorithmic comparisons also remain limited. Adaptive search has a larger computation budget than the deterministic baselines and retains their best result. No ablation attributes improvement specifically to edge reinforcement, and no optimality gap is available outside the small analytical example. The pooled pattern statistics describe the selected configuration matrix, including coincident and fixed-order families. They should not be interpreted as an unbiased distribution of all printable designs.')
para('Reduced external discard is a material-efficiency measure within one transition event. Additional motion, thermal demands and the consequences of combining polymers for recycling remain outside the analysis [2]. A future process assessment must include these effects together with calibrated structural acceptance. The current dataset provides reproducible computational evidence for allocation behaviour rather than a life-cycle or reinforcement claim.')
para('8 Conclusion','h1')
para('Transition reuse through a shared melt path was formulated as a composition-constrained scheduling problem and evaluated in 540 distinct configurations with 16,200 seeded adaptive runs. Conservation and the exact two-segment case establish that sufficient receiving capacity does not guarantee a feasible deposition order. All stored schedules satisfied the implemented volume and composition checks.')
para('In the 360 main configurations, adaptive search improved median utilisation over the better deterministic baseline in 34 cases, with zero median gain across the full matrix. The largest median gain was 47.73 percentage points for a concentric receiver. Fixed-order patterns showed no benefit, and the proposed staggered-finger pattern was not generally superior. Changing composition history and segment length affected allocation, while finer segments increased computation time without guaranteeing monotonic improvement.')
para('The supported outcome is selective improvement in modelled transition allocation. Stronger joints, physical printability and total process savings remain unverified. Calibration of discharge composition, overlap-aware geometry, a completion schedule and matched mechanical tests are required before the transition material can be accepted as a structural constituent.')
para('References','h1')
for p in refs:body.insert(len(body)-1,p)
for p in d.paragraphs:
 for r in p.runs:r.font.color.rgb=RGBColor(0,0,0)
 if re.match(r'^\d+ ',p.text):p.paragraph_format.keep_with_next=True;p.paragraph_format.first_line_indent=Cm(0)
 if re.match(r'^\d+\.\d+ ',p.text):p.paragraph_format.keep_with_next=True
 if p.text.startswith('Table '):p.paragraph_format.keep_with_next=True
 if re.match(r'^(Table|Figure) \d+\.',p.text):p.alignment=WD_ALIGN_PARAGRAPH.CENTER
# Format inline symbols and powers as editable Word runs, avoiding code-like caret notation.
token=re.compile(r'\b(?:Vcap|Velig|Vuse|Vdis|Vtr|Lref|Vi|Ai|li|si|ti|s0|ej|yj)\b|ΔVj|10\^-\d+|\^\(1/γ\)|\^γ')
for p in d.paragraphs:
 if p._p.xpath('.//m:oMath') or p.text.startswith('['):continue
 text=p.text
 if not token.search(text):continue
 p.clear();last=0
 for m in token.finditer(text):
  p.add_run(text[last:m.start()]);value=m.group()
  if value.startswith('10^'):
   p.add_run('10');r=p.add_run(value[3:]);r.font.superscript=True
  elif value.startswith('^'):
   r=p.add_run(value[1:]);r.font.superscript=True;r.italic=True
  else:
   base='ΔV' if value=='ΔVj' else value[0];suffix=value[len(base):]
   r=p.add_run(base);r.italic=True;r=p.add_run(suffix);r.font.subscript=True;r.italic=suffix in ('i','j')
  last=m.end()
 p.add_run(text[last:])
 for r in p.runs:r.font.name='Times New Roman';r.font.size=Pt(11);r.font.color.rgb=RGBColor(0,0,0)
for pr in d.element.xpath('.//w:pPr'):
 for border in list(pr.findall(qn('w:pBdr'))):pr.remove(border)
for t in d.element.xpath('.//w:t'):
 if t.text:t.text=t.text.replace('\u2013','-').replace('\u2014','-')
for t in d.tables:
 for row in t.rows:
  pr=row._tr.get_or_add_trPr()
  if pr.find(qn('w:cantSplit')) is None:pr.append(OxmlElement('w:cantSplit'))
d.core_properties.title='Transition material allocation and toolpath optimisation for interlocking hybrid FDM parts'
d.core_properties.subject='Computational allocation study with reproducible schedules and sensitivity analyses'
d.core_properties.author='Sven Maricic; Mihael Holi; Tin Nanndo Jovanovic'
# Author metadata supplied by the authors. Keep affiliations and ORCIDs distinct.
p=d.paragraphs[1];p.clear()
def author_run(text,bold=True,sup=False,size=11):
 r=p.add_run(text);r.font.name='Times New Roman';r.font.size=Pt(size);r.font.color.rgb=RGBColor(0,0,0);r.bold=bold;r.font.superscript=sup;return r
author_run('Sven Maricic');author_run('1,2',sup=True)
r=author_run('*',sup=True)
anchor=OxmlElement('w:footnoteReference');anchor.set(qn('w:id'),'1');anchor.set(qn('w:customMarkFollows'),'1');r._r.insert(1,anchor)
author_run(' [0000-0001-6644-3216]',False,size=10)
author_run(', Mihael Holi');author_run('1',sup=True);author_run(' [0009-0007-9960-6879]',False,size=10)
author_run(',').add_break()
author_run('Tin Nanndo Jovanovic');author_run('1',sup=True);author_run(' [0009-0004-6633-4174]',False,size=10)
d.save(out)
# Native first-page footnote; use the footnotes relationship already in the supplied template.
from zipfile import ZipFile
from lxml import etree as E
with ZipFile(out) as z:entries=[(deepcopy(i),z.read(i.filename)) for i in z.infolist()]
ns={'w':'http://schemas.openxmlformats.org/wordprocessingml/2006/main'};w='{'+ns['w']+'}'
foot=E.fromstring(next(data for info,data in entries if info.filename=='word/footnotes.xml'))
for existing in foot.findall('w:footnote',ns):
 if existing.get(w+'id')=='1':foot.remove(existing)
note=E.SubElement(foot,w+'footnote');note.set(w+'id','1');fp=E.SubElement(note,w+'p');pr=E.SubElement(fp,w+'pPr')
sp=E.SubElement(pr,w+'spacing');sp.set(w+'before','0');sp.set(w+'after','0');sp.set(w+'line','240');sp.set(w+'lineRule','auto')
ind=E.SubElement(pr,w+'ind');ind.set(w+'firstLine','0')
for text,sup in [('*',True),(' Corresponding author: Sven Maricic, e-mail: smaricic@unipu.hr',False)]:
 r=E.SubElement(fp,w+'r');rp=E.SubElement(r,w+'rPr');font=E.SubElement(rp,w+'rFonts');font.set(w+'ascii','Times New Roman');font.set(w+'hAnsi','Times New Roman')
 sz=E.SubElement(rp,w+'sz');sz.set(w+'val','20')
 if sup:E.SubElement(rp,w+'vertAlign').set(w+'val','superscript')
 t=E.SubElement(r,w+'t');t.set('{http://www.w3.org/XML/1998/namespace}space','preserve');t.text=text
with ZipFile(out,'w') as z:
 for info,data in entries:z.writestr(info,E.tostring(foot,encoding='UTF-8',xml_declaration=True,standalone=True) if info.filename=='word/footnotes.xml' else data)
check=Document(out);text='\n'.join(p.text for p in check.paragraphs)
assert '\u2013' not in text and '\u2014' not in text
assert len([p for p in check.paragraphs if re.match(r'^\[\d+\]',p.text)])==13
assert '34 cases' in text and '16,200' in text
print(out);print('Abstract words',len(abstract.split()),'Body words',len(text.split()),'Tables',len(check.tables),'Figures',len(check.inline_shapes),'Math nodes',len(check.element.xpath('.//m:oMath')))
